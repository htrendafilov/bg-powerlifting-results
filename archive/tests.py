import re
from collections import Counter
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.test import TestCase, override_settings

from archive.models import (
    AgeGroup, Athlete, AthletePhoto, Competition, Equipment, Event, Lift, MeetLevel,
    CompetitionFile, FileKind, Record, RecordOrigin, Result, Sex, SiteSettings,
)
from archive.services.commit import apply_import, prepare_rows
from archive.services.importers import _equipment_in_text, parse_upload
from archive.services.merge import duplicate_candidates, merge_athletes
from archive.services.visibility import visible_athletes, visible_competitions, visible_results
from archive.services.weight_classes import weight_class_for
from archive.services.names import (
    athlete_name_key, normalize_name, reverse_transliterate, transliterate,
)
from archive.services.records import recalculate_records


class MissingProtocolTests(TestCase):
    def test_past_meet_without_file_says_there_is_no_protocol(self):
        Competition.objects.create(
            name="Без протокол",
            start_date=date(2023, 5, 27),
            city="София",
            slug="test-missing",
        )
        listing = self.client.get("/competitions/")
        self.assertContains(listing, "Няма протокол")
        detail = self.client.get("/competitions/test-missing/")
        self.assertContains(detail, "Няма протокол")

    def test_future_meet_is_marked_as_not_held_yet(self):
        Competition.objects.create(name="Предстои", start_date=date(2026, 11, 1), slug="test-future")
        detail = self.client.get("/competitions/test-future/")
        self.assertContains(detail, "Още не е проведено")
        self.assertNotContains(detail, "Няма протокол")


class NameTests(TestCase):
    def test_home_page(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)

    def test_transliteration_and_order(self):
        self.assertEqual(transliterate("Рая Андонова"), "Raya Andonova")
        self.assertEqual(
            athlete_name_key("Рая Андонова", ""),
            athlete_name_key("", "Andonova Raya"),
        )


class ImporterTests(TestCase):
    def test_openpowerlifting_sheet(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Federation", "Date", "MeetCountry", "MeetState", "MeetTown", "MeetName", "Formula"])
        sheet.append(["БФСТ", "2026-04-03", "България", "", "Свищов", "Първи кръг", "IPF GL Points"])
        sheet.append([])
        sheet.append(
            [
                "Place",
                "Name",
                "Sex",
                "Country",
                "Division",
                "BodyweightKg",
                "WeightClassKg",
                "Squat1Kg",
                "Squat2Kg",
                "Squat3Kg",
                "Best3SquatKg",
                "Bench1Kg",
                "Bench2Kg",
                "Bench3Kg",
                "Best3BenchKg",
                "Deadlift1Kg",
                "Deadlift2Kg",
                "Deadlift3Kg",
                "Best3DeadliftKg",
                "TotalKg",
                "Points",
                "Event",
                "",
            ]
        )
        sheet.append([])
        sheet.append(
            [
                1,
                "Рая Андонова",
                "F",
                "BUL",
                "SJR",
                61.9,
                63,
                80,
                90,
                100,
                100,
                35,
                42.5,
                -45,
                42.5,
                105,
                112.5,
                122.5,
                122.5,
                265,
                58.65,
                "SBD",
                "НСА",
            ]
        )
        parsed = parse_upload("svishov.xlsx", _bytes(workbook))
        self.assertEqual(parsed.kind, "opl")
        self.assertEqual(parsed.meet_date, date(2026, 4, 3))
        self.assertEqual(parsed.city, "Свищов")
        self.assertEqual(len(parsed.rows), 1)
        row = parsed.rows[0]
        self.assertEqual(row.raw_name, "Рая Андонова")
        self.assertEqual(row.sex, Sex.F)
        self.assertEqual(row.age_group, AgeGroup.SUBJUNIOR)
        self.assertEqual(row.weight_class, "63")
        self.assertEqual(row.event, Event.SBD)
        self.assertEqual(row.best_squat, 100)
        self.assertEqual(row.attempts["bench3"], -45)
        self.assertEqual(row.best_bench, 42.5)
        self.assertEqual(row.total, 265)
        self.assertEqual(row.club, "НСА")
        self.assertEqual(row.equipment, "")

    def test_goodlift_bench_sheet(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Benchpress 2026 NSA • BG, Sofia • 30.05.2026"])
        sheet.append([])
        sheet.append([])
        sheet.append([])
        sheet.append(["PL.", "Name", "B.Date", "Nation", "Weight", "Lot", "1 Att.", "2 Att.", "3 Att.", "RESULT", "Pts."])
        sheet.append(["Subjuniors"])
        sheet.append(["- 66 kg"])
        sheet.append([1, "Karakashev Dimitar", "01.01.09", "Yunak", 66, 21, 85, 90, "X", 90, 72.631])
        parsed = parse_upload("bench.xlsx", _bytes(workbook))
        self.assertEqual(parsed.kind, "goodlift")
        self.assertEqual(parsed.meet_date, date(2026, 5, 30))
        self.assertEqual(len(parsed.rows), 1)
        row = parsed.rows[0]
        self.assertEqual(row.age_group, AgeGroup.SUBJUNIOR)
        self.assertEqual(row.weight_class, "66")
        self.assertEqual(row.event, Event.B)
        self.assertEqual(row.club, "Yunak")
        self.assertEqual(row.best_bench, 90)
        self.assertIsNone(row.attempts.get("bench3"))
        self.assertEqual(row.country, "")


class RecordTests(TestCase):
    def test_results_fill_empty_cells_and_keep_history(self):
        lifter = Athlete.objects.create(name_bg="Иван Петров", sex=Sex.M)
        other = Athlete.objects.create(name_bg="Петър Иванов", sex=Sex.M)
        Record.objects.create(
            sex=Sex.M,
            age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC,
            event=Event.SBD,
            lift=Lift.SQUAT,
            weight_class="93",
            value_kg=300,
            origin=RecordOrigin.STANDARD,
        )
        seed_holder = Athlete.objects.create(name_bg="Стар Рекорд", sex=Sex.M)
        seed = Record.objects.create(
            sex=Sex.M,
            age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC,
            event=Event.SBD,
            lift=Lift.BENCH,
            weight_class="93",
            value_kg=200,
            athlete=seed_holder,
            origin=RecordOrigin.SEED,
            valid_from=date(2019, 1, 1),
            note="БФСТ",
        )
        below = self._result(lifter, date(2024, 5, 1), best_squat=290, best_bench=190)
        above = self._result(other, date(2024, 9, 1), best_squat=310, best_bench=210)
        recalculate_records()
        self.assertFalse(Record.objects.filter(origin=RecordOrigin.RESULT, lift=Lift.SQUAT, result=below).exists())
        squat = Record.objects.get(origin=RecordOrigin.RESULT, lift=Lift.SQUAT, age_group=AgeGroup.OPEN, valid_to=None)
        self.assertEqual(squat.value_kg, 310)
        self.assertEqual(squat.athlete, other)
        seed.refresh_from_db()
        self.assertEqual(seed.valid_to, date(2024, 9, 1))
        current_bench = Record.objects.get(origin=RecordOrigin.RESULT, lift=Lift.BENCH, valid_to=None)
        self.assertEqual(current_bench.value_kg, 210)
        recalculate_records()
        self.assertEqual(Record.objects.filter(origin=RecordOrigin.RESULT, lift=Lift.SQUAT).count(), 1)

    def test_junior_also_sets_open_and_full_meet_bench_stays_separate(self):
        lifter = Athlete.objects.create(name_bg="Млад Състезател", sex=Sex.M)
        self._result(lifter, date(2024, 6, 1), age_group=AgeGroup.JUNIOR, best_squat=180, best_bench=120, best_deadlift=200, total=500)
        recalculate_records()
        self.assertTrue(Record.objects.filter(age_group=AgeGroup.JUNIOR, lift=Lift.TOTAL, value_kg=500).exists())
        self.assertTrue(Record.objects.filter(age_group=AgeGroup.OPEN, lift=Lift.TOTAL, value_kg=500).exists())
        self.assertFalse(Record.objects.filter(event=Event.B).exists())

    def test_doping_foreign_and_earlier_conflict(self):
        lifter = Athlete.objects.create(name_bg="Българин", sex=Sex.M)
        Record.objects.create(
            sex=Sex.M,
            age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC,
            event=Event.SBD,
            lift=Lift.DEADLIFT,
            weight_class="93",
            value_kg=300,
            athlete=lifter,
            origin=RecordOrigin.SEED,
            valid_from=date(2024, 6, 1),
        )
        early = self._result(lifter, date(2023, 1, 1), best_deadlift=310)
        self._result(lifter, date(2024, 8, 1), best_deadlift=320, place="DD")
        guest = Athlete.objects.create(name_lat="Foreign Guest", sex=Sex.M)
        self._result(guest, date(2024, 8, 2), best_deadlift=400, country="USA")
        recalculate_records()
        early.refresh_from_db()
        self.assertIn("по-ранна дата", early.review_note)
        current = Record.objects.get(lift=Lift.DEADLIFT, valid_to=None, origin=RecordOrigin.SEED)
        self.assertEqual(current.value_kg, 300)

    def test_import_creates_athlete_once(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Place", "Name", "Sex", "Country", "Division", "WeightClassKg", "Best3BenchKg", "Event"])
        sheet.append([1, "Иван Петров", "M", "BUL", "Open", "93", 150, "B"])
        parsed = parse_upload("bench.csv", b"")
        parsed = parse_upload("one.xlsx", _bytes(workbook))
        competition = Competition.objects.create(name="Лег", start_date=date(2025, 5, 1))
        summary = apply_import(
            competition,
            parsed,
            default_sex="",
            default_equipment=Equipment.CLASSIC,
            default_event="auto",
            replace=True,
        )
        self.assertEqual(summary["created"], 1)
        self.assertEqual(Athlete.objects.count(), 1)
        self.assertEqual(Record.objects.get(event=Event.B, lift=Lift.BENCH).value_kg, 150)
        apply_import(
            competition,
            parsed,
            default_sex="",
            default_equipment=Equipment.CLASSIC,
            default_event="auto",
            replace=True,
        )
        self.assertEqual(Athlete.objects.count(), 1)
        self.assertEqual(Result.objects.count(), 1)

    def test_photos_stay_hidden_under_18(self):
        minor = Athlete.objects.create(name_bg="Юноша", sex=Sex.M, birth_year=date.today().year - 16)
        adult = Athlete.objects.create(name_bg="Възрастен", sex=Sex.M, birth_year=date.today().year - 30)
        unknown = Athlete.objects.create(name_bg="Без година", sex=Sex.M)
        confirmed = Athlete.objects.create(name_bg="Потвърден", sex=Sex.M, adult_confirmed=True)
        self.assertFalse(minor.allows_public_photos)
        self.assertTrue(adult.allows_public_photos)
        self.assertFalse(unknown.allows_public_photos)
        self.assertTrue(confirmed.allows_public_photos)

    def _result(self, athlete, day, **fields):
        competition = Competition.objects.create(name=f"Турнир {day.isoformat()}", start_date=day)
        defaults = {
            "sex": athlete.sex,
            "age_group": AgeGroup.OPEN,
            "equipment": Equipment.CLASSIC,
            "event": Event.SBD,
            "weight_class": "93",
            "country": "България",
            "raw_name": athlete.display_name,
        }
        defaults.update(fields)
        return Result.objects.create(competition=competition, athlete=athlete, **defaults)


def _bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class RealProtocolTests(TestCase):
    """The two protocol shapes the federation actually publishes. Synthetic
    sheets missed every defect these files carry."""

    def _fixture(self, name):
        return (Path(__file__).parent / "tests_data" / name).read_bytes()

    def test_goodlift_bench_drops_team_and_best_lifter_tables(self):
        parsed = parse_upload("bench.xlsx", self._fixture("goodlift_bench_2026.xlsx"))
        self.assertEqual(parsed.kind, "goodlift")
        self.assertEqual(len(parsed.rows), 85)
        self.assertEqual(parsed.skipped, 57)
        self.assertIn("Nation (points)", parsed.skipped_sections)
        names = {row.raw_name for row in parsed.rows}
        for club in ("MNG", "NSA-Sofia", "STRONG", "Marek-13"):
            self.assertNotIn(club, names)

    def test_goodlift_seniors_are_not_filed_as_juniors(self):
        parsed = parse_upload("bench.xlsx", self._fixture("goodlift_bench_2026.xlsx"))
        by_group = Counter(row.age_group for row in parsed.rows)
        self.assertEqual(by_group[AgeGroup.OPEN], 31)
        self.assertEqual(by_group[AgeGroup.JUNIOR], 20)

    def test_goodlift_nation_column_of_a_domestic_meet_is_the_club(self):
        parsed = parse_upload("bench.xlsx", self._fixture("goodlift_bench_2026.xlsx"))
        self.assertEqual({row.country for row in parsed.rows}, {""})
        self.assertEqual(parsed.rows[0].club, "Yunak")

    def test_goodlift_bench_needs_an_equipment_choice(self):
        parsed = parse_upload("bench.xlsx", self._fixture("goodlift_bench_2026.xlsx"))
        _, blocked = prepare_rows(parsed, default_sex=Sex.M, default_equipment="auto", default_event="auto")
        self.assertEqual(len(blocked), 85)
        self.assertIn("Няма екипировка", blocked[0]["problems"][0])
        ready, blocked = prepare_rows(
            parsed, default_sex=Sex.M, default_equipment=Equipment.CLASSIC, default_event="auto"
        )
        self.assertEqual((len(ready), len(blocked)), (85, 0))

    def test_opl_sbd_protocol_still_imports_whole(self):
        parsed = parse_upload("svishtov.xlsx", self._fixture("opl_sbd_2026.xlsx"))
        self.assertEqual(parsed.kind, "opl")
        self.assertEqual(len(parsed.rows), 146)
        self.assertEqual(parsed.skipped, 0)
        self.assertEqual(parsed.meet_date, date(2026, 4, 3))
        self.assertEqual(parsed.city, "Свищов")
        by_group = Counter(row.age_group for row in parsed.rows)
        self.assertEqual(by_group[AgeGroup.OPEN], 61)
        # "с екип" in the meet title, so the title supplies it and the operator need not.
        self.assertEqual({row.equipment for row in parsed.rows}, {Equipment.EQUIPPED})


class GoodliftFullPowerTests(TestCase):
    def _sheet(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Powerlifting 2026 • BG, Sofia • 19.09.2026"])
        sheet.append([])
        sheet.append([
            "PL.", "Name", "B.Date", "Nation", "Weight", "Lot",
            "1 Att.", "2 Att.", "3 Att.",
            "1 Att.", "2 Att.", "3 Att.",
            "1 Att.", "2 Att.", "3 Att.",
            "RESULT", "Pts.",
        ])
        sheet.append(["Seniors"])
        sheet.append(["- 93 kg"])
        sheet.append([1, "Ivanov Ivan", "01.01.95", "Levski", 92.5, 7,
                      250, 260, 270, 150, 160, "X", 280, 290, 300, 730, 90.1])
        return _bytes(workbook)

    def test_three_attempt_blocks_land_on_the_right_lifts(self):
        row = parse_upload("full.xlsx", self._sheet()).rows[0]
        self.assertEqual(row.event, Event.SBD)
        self.assertEqual(row.age_group, AgeGroup.OPEN)
        self.assertEqual(row.attempts["squat3"], 270)
        self.assertEqual(row.attempts["bench2"], 160)
        self.assertEqual(row.attempts["deadlift3"], 300)
        self.assertEqual((row.best_squat, row.best_bench, row.best_deadlift), (270, 160, 300))
        self.assertEqual(row.total, 730)

    def test_an_unknown_attempt_layout_is_refused_not_guessed(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["PL.", "Name", "1 Att.", "2 Att.", "3 Att.", "1 Att.", "2 Att.", "3 Att.", "RESULT"])
        sheet.append(["Open"])
        sheet.append(["- 93 kg"])
        sheet.append([1, "Ivanov Ivan", 100, 110, 120, 130, 140, 150, 270])
        row = parse_upload("odd.xlsx", _bytes(workbook)).rows[0]
        self.assertTrue(any("Неразпознат Goodlift лист" in problem for problem in row.errors))


class ForeignLifterTests(TestCase):
    def _euro_sheet(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["European Bench Press Championships • Kaunas • 12.05.2026"])
        sheet.append([])
        sheet.append(["PL.", "Name", "B.Date", "Nation", "Weight", "Lot", "1 Att.", "2 Att.", "3 Att.", "RESULT", "Pts."])
        sheet.append(["Seniors"])
        sheet.append(["- 83 kg"])
        sheet.append([1, "Petrauskas Jonas", "01.01.95", "LTU", 82.5, 3, 200, 210, 220, 220, 130.0])
        sheet.append([2, "Kovacs Andras", "01.01.95", "HUN", 82.7, 5, 195, 205, 215, 215, 127.0])
        sheet.append([3, "Ivanov Ivan", "01.01.95", "BUL", 82.9, 4, 190, 200, "X", 200, 120.0])
        return _bytes(workbook)

    def test_nation_column_of_an_international_meet_is_the_country(self):
        parsed = parse_upload("euro.xlsx", self._euro_sheet())
        self.assertEqual([row.country for row in parsed.rows], ["LTU", "HUN", "BUL"])
        self.assertEqual({row.club for row in parsed.rows}, {""})

    def test_foreigners_do_not_take_bulgarian_records(self):
        parsed = parse_upload("euro.xlsx", self._euro_sheet())
        competition = Competition.objects.create(
            name="Европейско", start_date=date(2026, 5, 12), level=MeetLevel.INTERNATIONAL
        )
        apply_import(
            competition, parsed, default_sex=Sex.M, default_equipment=Equipment.CLASSIC,
            default_event="auto", replace=True,
        )
        self.assertEqual(Result.objects.count(), 3)
        holder = Record.objects.get(origin=RecordOrigin.RESULT, lift=Lift.BENCH, valid_to=None)
        self.assertEqual(holder.value_kg, 200)
        self.assertEqual(holder.athlete.name_lat, "Ivanov Ivan")

    def test_an_international_row_without_a_country_is_blocked(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Place", "Name", "Sex", "Division", "WeightClassKg", "Best3BenchKg", "Event"])
        sheet.append([1, "Ivanov Ivan", "M", "Open", "93", 150, "B"])
        parsed = parse_upload("nc.xlsx", _bytes(workbook))
        _, blocked = prepare_rows(
            parsed, default_sex="", default_equipment=Equipment.CLASSIC, default_event="auto",
            meet_level=MeetLevel.INTERNATIONAL,
        )
        self.assertIn("Няма държава", blocked[0]["problems"][0])


class MeetTitleEquipmentTests(TestCase):
    def test_combined_rounds_do_not_claim_to_be_raw(self):
        self.assertEqual(_equipment_in_text("2 кръг, вдигане от лег с и без екип"), "")
        self.assertEqual(_equipment_in_text("Силов трибой с и без екип"), "")
        self.assertEqual(_equipment_in_text("3 кръг, силов трибой без екип"), Equipment.CLASSIC)
        self.assertEqual(_equipment_in_text("Класически силов трибой"), Equipment.CLASSIC)
        self.assertEqual(_equipment_in_text("2 кръг, силов трибой с екип"), Equipment.EQUIPPED)
        self.assertEqual(_equipment_in_text("Екипировъчен силов трибой"), Equipment.EQUIPPED)

    def test_the_operator_overrides_the_title_but_not_a_column(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["MeetName", "Силов трибой без екип"])
        sheet.append([])
        sheet.append(["Place", "Name", "Sex", "Country", "Division", "WeightClassKg", "Best3BenchKg", "Event"])
        sheet.append([1, "Ivan Ivanov", "M", "BUL", "Open", "93", 150, "B"])
        parsed = parse_upload("t.xlsx", _bytes(workbook))
        self.assertEqual(parsed.rows[0].equipment_source, "title")
        ready, _ = prepare_rows(
            parsed, default_sex="", default_equipment=Equipment.EQUIPPED, default_event="auto"
        )
        self.assertEqual(ready[0].equipment, Equipment.EQUIPPED)


class UrlTests(TestCase):
    def test_english_paths_answer_and_old_ones_redirect(self):
        Competition.objects.create(name="Турнир", start_date=date(2025, 5, 1), slug="t")
        for path in ("/competitions/", "/records/", "/athletes/", "/competitions/t/"):
            self.assertEqual(self.client.get(path).status_code, 200, path)
        for old, new in (
            ("/sustezaniya/", "/competitions/"),
            ("/sustezaniya/t/", "/competitions/t/"),
            ("/rekordi/", "/records/"),
            ("/sastezateli/", "/athletes/"),
        ):
            response = self.client.get(old)
            self.assertEqual(response.status_code, 302, old)
            self.assertEqual(response["Location"], new)

    def test_redirect_keeps_the_query_string(self):
        response = self.client.get("/rekordi/?sex=F&event=B")
        self.assertEqual(response["Location"], "/records/?sex=F&event=B")


class OrderingTests(TestCase):
    def test_competitions_are_listed_newest_first_despite_the_annotation(self):
        for day, name, slug in (
            (date(2021, 6, 5), "стар", "old"),
            (date(2026, 4, 3), "нов", "new"),
            (date(2023, 7, 1), "среден", "mid"),
        ):
            Competition.objects.create(name=name, start_date=day, slug=slug)
        home = self.client.get("/")
        listing = self.client.get("/competitions/")
        for page in (home, listing):
            body = page.content.decode()
            order = [body.index(f"/competitions/{slug}/") for slug in ("new", "mid", "old")]
            self.assertEqual(order, sorted(order), page.request["PATH_INFO"])
        years = [int(y) for y in re.findall(r"<h2>(\d{4})</h2>", listing.content.decode())]
        self.assertEqual(years, sorted(years, reverse=True))


class DeadliftOnlyTests(TestCase):
    def test_a_deadlift_sheet_is_not_recorded_as_bench(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Национален шампионат по мъртва тяга • София • 26.11.2026"])
        sheet.append([])
        sheet.append(["PL.", "Name", "Nation", "Weight", "1 Att.", "2 Att.", "3 Att.", "RESULT", "Pts."])
        sheet.append(["Seniors"])
        sheet.append(["- 93 kg"])
        sheet.append([1, "Ivanov Ivan", "Levski", 92.5, 250, 265, 280, 280, 90.1])
        row = parse_upload("dl.xlsx", _bytes(workbook)).rows[0]
        self.assertEqual(row.event, Event.D)
        self.assertEqual(row.best_deadlift, 280)
        self.assertIsNone(row.best_bench)
        self.assertEqual(row.attempts["deadlift3"], 280)
        self.assertNotIn("bench1", row.attempts)

    def test_a_deadlift_result_never_becomes_a_record(self):
        athlete = Athlete.objects.create(name_bg="Тягаджия", sex=Sex.M)
        competition = Competition.objects.create(name="Тяга", start_date=date(2026, 11, 26))
        Result.objects.create(
            competition=competition, athlete=athlete, sex=Sex.M, age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC, event=Event.D, weight_class="93",
            country="България", best_deadlift=400,
        )
        recalculate_records()
        self.assertEqual(Record.objects.count(), 0)

    def test_a_meet_can_be_excluded_from_records_by_hand(self):
        athlete = Athlete.objects.create(name_bg="Показен", sex=Sex.M)
        competition = Competition.objects.create(name="Показен", start_date=date(2026, 3, 1))
        result = Result.objects.create(
            competition=competition, athlete=athlete, sex=Sex.M, age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC, event=Event.SBD, weight_class="93",
            country="България", best_bench=250, counts_for_records=False,
        )
        recalculate_records()
        self.assertEqual(Record.objects.count(), 0)
        result.counts_for_records = True
        result.save()
        recalculate_records()
        self.assertEqual(Record.objects.filter(lift=Lift.BENCH, value_kg=250).count(), 1)


@override_settings(
    DEBUG=False, SECURE_SSL_REDIRECT=True, SECURE_HSTS_SECONDS=31536000,
    SECURE_HSTS_INCLUDE_SUBDOMAINS=True, SECURE_PROXY_SSL_HEADER=("HTTP_X_FORWARDED_PROTO", "https"),
)
class HttpsTests(TestCase):
    def test_plain_http_is_redirected(self):
        response = self.client.get("/competitions/")
        self.assertEqual(response.status_code, 301)
        self.assertTrue(response["Location"].startswith("https://"))

    def test_a_request_behind_the_tunnel_is_served_with_hsts(self):
        response = self.client.get("/competitions/", HTTP_X_FORWARDED_PROTO="https")
        self.assertEqual(response.status_code, 200)
        self.assertIn("max-age=31536000", response["Strict-Transport-Security"])
        self.assertIn("includeSubDomains", response["Strict-Transport-Security"])


class SectionHeadingTests(TestCase):
    """A sheet whose rows carry their own Division must not lose them to a
    heading the label parser cannot read (real file: Кърджали 2023)."""

    def test_decorative_headings_do_not_swallow_the_rows(self):
        data = (Path(__file__).parent / "tests_data" / "opl_sections_2023.xlsx").read_bytes()
        parsed = parse_upload("k.xlsx", data)
        self.assertEqual(len(parsed.rows), 87)
        self.assertEqual(parsed.skipped, 0)

    def test_a_context_dependent_sheet_still_drops_unknown_blocks(self):
        data = (Path(__file__).parent / "tests_data" / "goodlift_bench_2026.xlsx").read_bytes()
        parsed = parse_upload("b.xlsx", data)
        self.assertEqual(parsed.skipped, 57)


class AgeCodeTests(TestCase):
    """The federation heads one column "ГР" but fills it two different ways:
    the division's lowest age, or an ordinal code (real files: Дупница 2023
    against Горна Оряховица 2022)."""

    HEAD = ["\u2116", "\u0418\u043c\u0435", "\u0424\u0430\u043c\u0438\u043b\u0438\u044f",
            "\u041e\u0442\u0431\u043e\u0440", "\u0413\u0420", "\u0442\u0435\u0433\u043b\u043e",
            "\u043a\u0430\u0442.", "\u043b\u0435\u04331", "\u043b\u0435\u04332", "\u043b\u0435\u04333"]

    def _parse(self, codes):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(self.HEAD)
        for index, code in enumerate(codes, start=1):
            sheet.append([str(index), f"\u0418\u043c\u0435{index}", f"\u0424\u0430\u043c{index}",
                          "\u041d\u0421\u0410", str(code), "74.5", "83", "100", "110", "120"])
        buffer = BytesIO()
        workbook.save(buffer)
        return parse_upload("gr.xlsx", buffer.getvalue())

    def test_ordinal_codes_become_divisions(self):
        parsed = self._parse([0, 1, 2, 3, 4, 18, 23])
        self.assertEqual(
            [row.age_group for row in parsed.rows],
            [AgeGroup.OPEN, AgeGroup.M1, AgeGroup.M2, AgeGroup.M3, AgeGroup.M4,
             AgeGroup.SUBJUNIOR, AgeGroup.JUNIOR],
        )

    def test_a_column_of_ages_is_left_for_the_age_bands(self):
        parsed = self._parse([18, 23, 24, 40, 70])
        self.assertEqual([row.age_group for row in parsed.rows], [""] * 5)
        self.assertEqual([int(row.age) for row in parsed.rows], [18, 23, 24, 40, 70])


class HeaderWithoutNameTests(TestCase):
    """A header that ranks and measures but never names the lifter used to
    crash: the place column lives in layout.places, never in layout.mapping."""

    def test_a_nameless_header_is_refused_instead_of_raising(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["\u2116", "\u0442\u0435\u0433\u043b\u043e", "\u043e\u0442\u0431\u043e\u0440",
                      "\u043f\u043e\u043b", "\u0442\u043e\u0447\u043a\u0438"])
        sheet.append(["1", "74.5", "\u041d\u0421\u0410", "M", "88.1"])
        buffer = BytesIO()
        workbook.save(buffer)
        parsed = parse_upload("nameless.xlsx", buffer.getvalue())
        self.assertEqual(parsed.rows, [])
        self.assertTrue(parsed.errors)


class BulgarianProtocolTests(TestCase):
    """The shape the federation publishes most often: a Bulgarian header row."""

    def _fixture(self, name):
        return (Path(__file__).parent / "tests_data" / name).read_bytes()

    def test_bulgarian_headers_are_read(self):
        parsed = parse_upload("sofia.xlsx", self._fixture("bg_sbd_2024.xlsx"))
        self.assertEqual(len(parsed.rows), 170)
        ready, blocked = prepare_rows(
            parsed, default_sex="", default_equipment="auto", default_event="auto"
        )
        self.assertEqual((len(ready), len(blocked)), (170, 0))
        row = parsed.rows[0]
        self.assertTrue(row.raw_name)
        self.assertTrue(row.club)
        self.assertEqual(row.event, Event.SBD)

    def test_a_combined_division_code_supplies_sex_equipment_and_event(self):
        parsed = parse_upload("sofia.xlsx", self._fixture("bg_sbd_2024.xlsx"))
        by_sex = Counter(row.sex for row in parsed.rows)
        self.assertEqual(by_sex[Sex.F], 34)
        self.assertEqual(by_sex[Sex.M], 136)
        self.assertEqual({row.equipment for row in parsed.rows}, {Equipment.CLASSIC})
        self.assertEqual({row.equipment_source for row in parsed.rows}, {"row"})

    def test_a_deadlift_meet_is_imported_but_sets_no_record(self):
        parsed = parse_upload("dl.xlsx", self._fixture("bg_deadlift_2023.xlsx"))
        self.assertEqual({row.event for row in parsed.rows}, {Event.D})
        ready, _ = prepare_rows(
            parsed, default_sex="", default_equipment="auto", default_event="auto"
        )
        competition = Competition.objects.create(
            name="Мъртва тяга", start_date=date(2023, 11, 26), slug="dl-2023"
        )
        summary = apply_import(
            competition, parsed, default_sex="", default_equipment="auto",
            default_event="auto", replace=True,
        )
        self.assertEqual(summary["created"], 96)
        self.assertEqual(len(ready), 96)
        self.assertEqual(Record.objects.count(), 0)


class MergedDivisionCodeTests(TestCase):
    def test_sex_prefixed_divisions_are_understood(self):
        from archive.services.importers import _division_label

        self.assertEqual(_division_label("F-Jr"), AgeGroup.JUNIOR)
        self.assertEqual(_division_label("M-O"), AgeGroup.OPEN)
        self.assertEqual(_division_label("M-T3"), AgeGroup.SUBJUNIOR)
        self.assertEqual(_division_label("M-M1"), AgeGroup.M1)
        self.assertEqual(_division_label("M1"), AgeGroup.M1)


class CompetitionLayoutTests(TestCase):
    def _result(self, competition, name, sex, age_group, weight_class, place, equipment=Equipment.CLASSIC):
        athlete = Athlete.objects.create(name_bg=name, sex=sex)
        return Result.objects.create(
            competition=competition, athlete=athlete, raw_name=name, sex=sex,
            age_group=age_group, equipment=equipment, event=Event.SBD,
            weight_class=weight_class, place=place, country="България", best_bench=100,
        )

    def setUp(self):
        self.competition = Competition.objects.create(
            name="Турнир", start_date=date(2026, 4, 3), slug="layout"
        )
        self._result(self.competition, "Мъж 120", Sex.M, AgeGroup.OPEN, "120", "1")
        self._result(self.competition, "Мъж 83", Sex.M, AgeGroup.OPEN, "83", "1")
        self._result(self.competition, "Жена 84+", Sex.F, AgeGroup.OPEN, "84+", "1")
        self._result(self.competition, "Жена 84", Sex.F, AgeGroup.OPEN, "84", "1")
        self._result(self.competition, "Жена 63 юн", Sex.F, AgeGroup.SUBJUNIOR, "63", "2")
        self._result(self.competition, "Жена 63 отк", Sex.F, AgeGroup.OPEN, "63", "1")

    def test_women_come_first_then_classes_ascending(self):
        from archive.views import _group_results

        groups = _group_results(list(self.competition.results.select_related("athlete")))
        self.assertEqual([g["label"] for g in groups], ["Жени", "Мъже"])
        self.assertEqual([c["weight_class"] for c in groups[0]["classes"]], ["63", "84", "84+"])
        self.assertEqual([c["weight_class"] for c in groups[1]["classes"]], ["83", "120"])

    def test_a_class_keeps_its_divisions_together(self):
        from archive.views import _group_results

        groups = _group_results(list(self.competition.results.select_related("athlete")))
        names = [r.raw_name for r in groups[0]["classes"][0]["results"]]
        self.assertEqual(names, ["Жена 63 юн", "Жена 63 отк"])

    def test_the_page_separates_equipment_from_the_weight_class(self):
        page = self.client.get("/competitions/layout/").content.decode()
        self.assertIn(">Екип.<", page)
        self.assertNotIn("120 кл.", page)
        self.assertNotIn("120 екип", page)
        self.assertIn("120 кг", page)
        self.assertIn('class="sex-heading"><span>Жени</span>', page)


class MergeAthleteTests(TestCase):
    def setUp(self):
        self.keep = Athlete.objects.create(name_bg="Роберт Михайлов", sex=Sex.M, birth_year=1990)
        self.dupe = Athlete.objects.create(name_lat="Robert Michailov", sex=Sex.M)
        self.competition = Competition.objects.create(
            name="Европейско", start_date=date(2013, 5, 7), slug="euro-2013",
            level=MeetLevel.INTERNATIONAL,
        )
        self.result = Result.objects.create(
            competition=self.competition, athlete=self.dupe, raw_name="Robert Michailov",
            sex=Sex.M, age_group=AgeGroup.OPEN, equipment=Equipment.EQUIPPED, event=Event.SBD,
            weight_class="83", country="България", best_bench=222.5,
        )
        AthletePhoto.objects.create(athlete=self.dupe, year=2013)

    def test_everything_moves_to_the_survivor(self):
        slug = self.keep.slug
        merge_athletes(self.keep, [self.dupe])
        self.keep.refresh_from_db()
        self.assertEqual(self.keep.results.count(), 1)
        self.assertEqual(self.keep.photos.count(), 1)
        self.assertFalse(Athlete.objects.filter(pk=self.dupe.pk).exists())
        self.assertEqual(self.keep.slug, slug)
        self.assertIn("Слети: Robert Michailov", self.keep.notes)

    def test_the_survivor_keeps_what_it_has_and_gains_what_it_lacks(self):
        self.dupe.birth_year = 1991
        self.dupe.adult_confirmed = True
        self.dupe.save()
        merge_athletes(self.keep, [self.dupe])
        self.keep.refresh_from_db()
        self.assertEqual(self.keep.birth_year, 1990)
        self.assertEqual(self.keep.name_lat, "Robert Mihaylov")
        self.assertTrue(self.keep.adult_confirmed)

    def test_records_follow_the_results(self):
        recalculate_records()
        self.assertEqual(Record.objects.get(lift=Lift.BENCH).athlete, self.dupe)
        merge_athletes(self.keep, [self.dupe])
        self.assertEqual(Record.objects.get(lift=Lift.BENCH).athlete, self.keep)

    def test_a_seeded_record_is_carried_over_too(self):
        seed = Record.objects.create(
            sex=Sex.M, age_group=AgeGroup.OPEN, equipment=Equipment.EQUIPPED, event=Event.SBD,
            lift=Lift.SQUAT, weight_class="83", value_kg=300, athlete=self.dupe,
            origin=RecordOrigin.SEED, valid_from=date(2010, 1, 1),
        )
        merge_athletes(self.keep, [self.dupe])
        seed.refresh_from_db()
        self.assertEqual(seed.athlete, self.keep)

    def test_candidates_find_the_spelling_variants(self):
        pairs = duplicate_candidates()
        found = {(one.pk, two.pk) for _, one, two in pairs}
        self.assertIn(tuple(sorted((self.keep.pk, self.dupe.pk))), found)

    def test_candidates_pair_a_two_part_name_with_a_three_part_one(self):
        Result.objects.all().delete()
        Athlete.objects.all().delete()
        short = Athlete.objects.create(name_lat="Ivan Ivanov", sex=Sex.M)
        full = Athlete.objects.create(name_bg="Иван Петров Иванов", sex=Sex.M)
        found = {tuple(sorted((one.pk, two.pk))) for _, one, two in duplicate_candidates()}
        self.assertIn(tuple(sorted((short.pk, full.pk))), found)

    def test_a_different_person_with_a_different_name_is_not_paired(self):
        Result.objects.all().delete()
        Athlete.objects.all().delete()
        one = Athlete.objects.create(name_lat="Ivan Ivanov", sex=Sex.M)
        two = Athlete.objects.create(name_lat="Georgi Stoev", sex=Sex.M)
        found = {tuple(sorted((a.pk, b.pk))) for _, a, b in duplicate_candidates()}
        self.assertNotIn(tuple(sorted((one.pk, two.pk))), found)


class MergeAdminTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User

        self.admin = User.objects.create_superuser("admin", "a@b.bg", "pw-for-tests-only")
        self.client.force_login(self.admin)
        self.keep = Athlete.objects.create(name_bg="Роберт Михайлов", sex=Sex.M)
        self.dupe = Athlete.objects.create(name_lat="Robert Michailov", sex=Sex.M)

    def test_the_action_asks_before_it_merges(self):
        response = self.client.post(
            "/admin/archive/athlete/",
            {"action": "merge_selected", "_selected_action": [self.keep.pk, self.dupe.pk]},
        )
        self.assertContains(response, "Сливане на състезатели")
        self.assertEqual(Athlete.objects.count(), 2)

    def test_confirming_merges_into_the_chosen_record(self):
        response = self.client.post(
            "/admin/archive/athlete/",
            {
                "action": "merge_selected",
                "_selected_action": [self.keep.pk, self.dupe.pk],
                "confirm_merge": "1",
                "target": str(self.keep.pk),
            },
            follow=True,
        )
        self.assertEqual(Athlete.objects.count(), 1)
        self.assertEqual(Athlete.objects.get().pk, self.keep.pk)
        self.assertContains(response, "слети 1")

    def test_one_selected_athlete_is_refused(self):
        self.client.post(
            "/admin/archive/athlete/",
            {"action": "merge_selected", "_selected_action": [self.keep.pk]},
        )
        self.assertEqual(Athlete.objects.count(), 2)

    def test_the_duplicates_page_lists_the_pair(self):
        response = self.client.get("/admin/archive/athlete/duplicates/")
        self.assertContains(response, "Роберт Михайлов")
        self.assertContains(response, "Robert Michailov")


class BulgarianNameTests(TestCase):
    def test_a_standard_transliteration_comes_back_unchanged(self):
        for name in ("Иван Петров", "Рая Андонова", "Калоян Панайотов",
                     "Мария Табакова-Трендафилова", "Адриана Райкова"):
            self.assertEqual(reverse_transliterate(transliterate(name)), name)

    def test_the_spellings_the_sources_actually_use(self):
        cases = {
            "Christo Christov": "Христо Христов",
            "Michailov": "Михайлов",
            "Alexander": "Александър",
            "Ivailo": "Ивайло",
            "Georgy": "Георги",
            "Evgeniy": "Евгени",
            "Iordan": "Йордан",
            "Ilia": "Илия",
        }
        for latin, expected in cases.items():
            got = " ".join(reverse_transliterate(part) for part in latin.split())
            self.assertEqual(got, expected, latin)

    def test_the_command_prefers_a_name_read_from_a_bulgarian_protocol(self):
        Athlete.objects.create(name_bg="Кольо Иванов", sex=Sex.M)
        latin_only = Athlete.objects.create(name_lat="Koljo Ivanov", sex=Sex.M)
        call_command("bulgarize_names", verbosity=0)
        latin_only.refresh_from_db()
        self.assertEqual(latin_only.name_bg, "Кольо Иванов")
        self.assertTrue(latin_only.name_bg_auto)

    def test_a_name_read_from_a_protocol_is_never_overwritten(self):
        athlete = Athlete.objects.create(name_bg="Иван Петров", name_lat="Ivan Petrov", sex=Sex.M)
        call_command("bulgarize_names", verbosity=0)
        athlete.refresh_from_db()
        self.assertEqual(athlete.name_bg, "Иван Петров")
        self.assertFalse(athlete.name_bg_auto)

    def test_the_disambiguator_openpowerlifting_adds_is_kept(self):
        athlete = Athlete.objects.create(name_lat="Alexander Pavlov #2", sex=Sex.M)
        call_command("bulgarize_names", verbosity=0)
        athlete.refresh_from_db()
        self.assertEqual(athlete.name_bg, "Александър Павлов #2")

    def test_derived_names_are_flagged_for_review(self):
        Athlete.objects.create(name_lat="Ivaylo Hristov", sex=Sex.M)
        call_command("bulgarize_names", verbosity=0)
        self.assertEqual(Athlete.objects.filter(name_bg_auto=True).count(), 1)


class OldWeightClassTests(TestCase):
    def setUp(self):
        self.old_meet = Competition.objects.create(
            name="Старо", start_date=date(2005, 5, 1), slug="staro"
        )
        self.new_meet = Competition.objects.create(
            name="Ново", start_date=date(2025, 5, 1), slug="novo"
        )
        self.empty_meet = Competition.objects.create(
            name="Без протокол", start_date=date(2024, 5, 1), slug="bez"
        )
        self.old_lifter = self._result(self.old_meet, "Стар", Sex.M, "82.5")
        self.new_lifter = self._result(self.new_meet, "Нов", Sex.M, "83")
        self.both = self._result(self.new_meet, "И двете", Sex.M, "93")
        self._result(self.old_meet, "И двете", Sex.M, "90", athlete=self.both.athlete)

    def _result(self, competition, name, sex, weight_class, athlete=None):
        athlete = athlete or Athlete.objects.create(name_bg=name, sex=sex)
        return Result.objects.create(
            competition=competition, athlete=athlete, raw_name=name, sex=sex,
            age_group=AgeGroup.OPEN, equipment=Equipment.CLASSIC, event=Event.SBD,
            weight_class=weight_class, country="България", best_bench=100,
        )

    def _show(self, value):
        settings_row = SiteSettings.load()
        settings_row.show_old_weight_classes = value
        settings_row.save()

    def test_old_classes_are_hidden_by_default(self):
        self.assertFalse(SiteSettings.load().show_old_weight_classes)
        self.assertEqual(visible_results().count(), 2)

    def test_a_meet_held_only_in_old_classes_disappears(self):
        visible = visible_competitions(Competition.objects.all())
        self.assertNotIn(self.old_meet, visible)
        self.assertIn(self.new_meet, visible)
        self.assertEqual(self.client.get("/competitions/staro/").status_code, 404)

    def test_a_meet_with_no_results_at_all_stays(self):
        self.assertIn(self.empty_meet, visible_competitions(Competition.objects.all()))

    def test_a_lifter_seen_only_in_old_classes_disappears(self):
        visible = visible_athletes(Athlete.objects.all())
        self.assertNotIn(self.old_lifter.athlete, visible)
        self.assertIn(self.new_lifter.athlete, visible)
        self.assertEqual(self.client.get(f"/athletes/{self.old_lifter.athlete.slug}/").status_code, 404)

    def test_a_lifter_who_spans_both_keeps_only_the_current_starts(self):
        page = self.client.get(f"/athletes/{self.both.athlete.slug}/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Ново")
        self.assertNotContains(page, "Старо")

    def test_the_records_table_drops_the_old_classes(self):
        recalculate_records()
        page = self.client.get("/records/?sex=M&age_group=open&equipment=classic&event=SBD")
        self.assertNotContains(page, "<td>82.5</td>", html=False)
        self.assertContains(page, "83")

    def test_turning_the_setting_on_brings_everything_back(self):
        self._show(True)
        self.assertEqual(visible_results().count(), 4)
        self.assertIn(self.old_meet, visible_competitions(Competition.objects.all()))
        self.assertIn(self.old_lifter.athlete, visible_athletes(Athlete.objects.all()))
        self.assertEqual(self.client.get("/competitions/staro/").status_code, 200)

    def test_a_womens_52_is_current_while_a_mens_52_is_not(self):
        woman = self._result(self.new_meet, "Жена", Sex.F, "52")
        man = self._result(self.new_meet, "Мъж", Sex.M, "52")
        visible = set(visible_results().values_list("pk", flat=True))
        self.assertIn(woman.pk, visible)
        self.assertNotIn(man.pk, visible)

    def test_a_row_without_a_weight_class_stays_visible(self):
        missed = self._result(self.new_meet, "Неявил се", Sex.M, "")
        self.assertIn(missed.pk, set(visible_results().values_list("pk", flat=True)))

    def test_only_one_settings_row_can_exist(self):
        SiteSettings.objects.create(show_old_weight_classes=True)
        SiteSettings.objects.create(show_old_weight_classes=False)
        self.assertEqual(SiteSettings.objects.count(), 1)


class CompetitionLevelFilterTests(TestCase):
    def setUp(self):
        Competition.objects.create(name="Републиканско", start_date=date(2025, 5, 1),
                                   slug="nat", level=MeetLevel.NATIONAL)
        Competition.objects.create(name="Европейско", start_date=date(2025, 6, 1),
                                   slug="int", level=MeetLevel.INTERNATIONAL)

    def test_the_filter_narrows_to_one_level(self):
        both = self.client.get("/competitions/")
        self.assertContains(both, "Републиканско")
        self.assertContains(both, "Европейско")

        national = self.client.get("/competitions/?level=national")
        self.assertContains(national, "Републиканско")
        self.assertNotContains(national, "Европейско")

        international = self.client.get("/competitions/?level=international")
        self.assertContains(international, "Европейско")
        self.assertNotContains(international, "Републиканско")

    def test_a_nonsense_level_shows_everything(self):
        page = self.client.get("/competitions/?level=zzz")
        self.assertContains(page, "Републиканско")
        self.assertContains(page, "Европейско")

    def test_the_chosen_level_stays_selected(self):
        page = self.client.get("/competitions/?level=international")
        self.assertContains(page, 'value="international" selected')


class MergeNamePreferenceTests(TestCase):
    def test_a_name_from_a_protocol_beats_a_converted_one(self):
        converted = Athlete.objects.create(
            name_bg="Роберт Михаилов", name_lat="Robert Mihailov", sex=Sex.M, name_bg_auto=True
        )
        from_protocol = Athlete.objects.create(
            name_bg="Роберт Михайлов", name_lat="Robert Mihaylov", sex=Sex.M
        )
        merge_athletes(converted, [from_protocol])
        converted.refresh_from_db()
        self.assertEqual(converted.name_bg, "Роберт Михайлов")
        self.assertFalse(converted.name_bg_auto)

    def test_a_protocol_name_is_not_replaced_by_a_converted_one(self):
        from_protocol = Athlete.objects.create(name_bg="Роберт Михайлов", sex=Sex.M)
        converted = Athlete.objects.create(
            name_bg="Роберт Михаилов", name_lat="Robert Mihailov", sex=Sex.M, name_bg_auto=True
        )
        merge_athletes(from_protocol, [converted])
        from_protocol.refresh_from_db()
        self.assertEqual(from_protocol.name_bg, "Роберт Михайлов")
        self.assertFalse(from_protocol.name_bg_auto)


class WeightClassFromBodyweightTests(TestCase):
    def test_the_class_is_the_first_limit_the_lifter_makes(self):
        modern = date(2026, 1, 31)
        self.assertEqual(weight_class_for(Sex.M, Decimal("92.58"), modern), "93")
        self.assertEqual(weight_class_for(Sex.M, Decimal("93.00"), modern), "93")
        self.assertEqual(weight_class_for(Sex.M, Decimal("93.01"), modern), "105")
        self.assertEqual(weight_class_for(Sex.F, Decimal("63.25"), modern), "69")
        self.assertEqual(weight_class_for(Sex.F, Decimal("63.00"), modern), "63")

    def test_a_lifter_over_the_top_class_lands_in_the_plus(self):
        self.assertEqual(weight_class_for(Sex.M, Decimal("140"), date(2026, 1, 1)), "120+")
        self.assertEqual(weight_class_for(Sex.F, Decimal("95"), date(2026, 1, 1)), "84+")

    def test_a_result_from_before_2011_uses_the_old_classes(self):
        old = date(2000, 5, 20)
        self.assertEqual(weight_class_for(Sex.M, Decimal("149.4"), old), "125+")
        self.assertEqual(weight_class_for(Sex.M, Decimal("82.4"), old), "82.5")
        self.assertEqual(weight_class_for(Sex.F, Decimal("59"), old), "60")

    def test_the_lightest_class_belongs_to_the_youngest_groups(self):
        day = date(2026, 1, 1)
        self.assertEqual(weight_class_for(Sex.M, Decimal("52"), day, AgeGroup.SUBJUNIOR), "53")
        self.assertEqual(weight_class_for(Sex.M, Decimal("52"), day, AgeGroup.OPEN), "59")

    def test_nothing_is_invented_without_a_bodyweight(self):
        self.assertEqual(weight_class_for(Sex.M, None, date(2026, 1, 1)), "")

    def test_the_command_fills_the_gap_and_the_result_reaches_the_records(self):
        athlete = Athlete.objects.create(name_bg="Емил Кръстев", sex=Sex.M)
        competition = Competition.objects.create(
            name="Sheffield Powerlifting Championships", start_date=date(2026, 1, 31),
            slug="sheffield-2026", level=MeetLevel.INTERNATIONAL, city="Sheffield",
        )
        result = Result.objects.create(
            competition=competition, athlete=athlete, raw_name="Emil Krastev", sex=Sex.M,
            age_group=AgeGroup.OPEN, equipment=Equipment.CLASSIC, event=Event.SBD,
            weight_class="", bodyweight=Decimal("92.58"), country="България",
            best_squat=315, best_bench=Decimal("237.5"), best_deadlift=Decimal("367.5"), total=920,
        )
        recalculate_records()
        self.assertEqual(Record.objects.count(), 0)

        call_command("fill_weight_classes", verbosity=0)
        result.refresh_from_db()
        self.assertEqual(result.weight_class, "93")
        self.assertEqual(
            Record.objects.get(lift=Lift.TOTAL, weight_class="93", valid_to=None).value_kg, 920
        )

    def test_a_class_already_on_the_row_is_left_alone(self):
        athlete = Athlete.objects.create(name_bg="Тест", sex=Sex.M)
        competition = Competition.objects.create(name="T", start_date=date(2026, 1, 1), slug="t")
        result = Result.objects.create(
            competition=competition, athlete=athlete, sex=Sex.M, age_group=AgeGroup.OPEN,
            equipment=Equipment.CLASSIC, event=Event.SBD, weight_class="105",
            bodyweight=Decimal("92.5"), country="България", best_bench=100,
        )
        call_command("fill_weight_classes", verbosity=0)
        result.refresh_from_db()
        self.assertEqual(result.weight_class, "105")


class NonScoringPlaceTests(TestCase):
    def _result(self, place, **fields):
        athlete = Athlete.objects.create(name_bg=f"Лифтьор {place}", sex=Sex.M)
        competition = Competition.objects.create(
            name=f"Турнир {place}", start_date=date(2025, 3, 8), slug=f"t-{place.lower()}"
        )
        return Result.objects.create(
            competition=competition, athlete=athlete, raw_name=athlete.name_bg, sex=Sex.M,
            age_group=AgeGroup.OPEN, equipment=Equipment.CLASSIC, event=Event.SBD,
            weight_class="105", country="България", place=place, **fields,
        )

    def test_a_guest_start_sets_no_record(self):
        guest = self._result("G", best_squat=310, best_deadlift=Decimal("357.5"), total=900)
        self.assertFalse(guest.counts_for_bulgarian_records)
        recalculate_records()
        self.assertEqual(Record.objects.count(), 0)

    def test_a_disqualified_start_sets_no_record_even_with_good_lifts(self):
        dq = self._result("DQ", best_squat=350, best_bench=225)
        self.assertFalse(dq.counts_for_bulgarian_records)
        recalculate_records()
        self.assertEqual(Record.objects.count(), 0)

    def test_doping_and_no_show_still_set_nothing(self):
        for place in ("DD", "NS"):
            self.assertFalse(self._result(place, best_squat=400).counts_for_bulgarian_records)

    def test_a_placed_start_still_sets_records(self):
        placed = self._result("3", best_squat=300, best_bench=200, best_deadlift=300, total=800)
        self.assertTrue(placed.counts_for_bulgarian_records)
        recalculate_records()
        self.assertEqual(Record.objects.filter(lift=Lift.TOTAL, value_kg=800).count(), 1)

    def test_a_guest_does_not_displace_a_placed_lifter(self):
        self._result("3", best_squat=300, total=800)
        self._result("G", best_squat=350, total=900)
        recalculate_records()
        squat = Record.objects.get(lift=Lift.SQUAT, valid_to=None)
        self.assertEqual(squat.value_kg, 300)

    def test_dsq_is_stored_as_dq(self):
        from archive.services.importers import _clean_place

        self.assertEqual(_clean_place("DSQ"), "DQ")
        self.assertEqual(_clean_place("DNS"), "NS")
        self.assertEqual(_clean_place("G"), "G")


class DerivedRowFieldsTests(TestCase):
    def _sheet(self, header, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        return _bytes(workbook)

    def test_the_division_codes_the_protocols_use(self):
        from archive.services.importers import _division_label

        for code, expected in (
            ("F-Sj", AgeGroup.SUBJUNIOR), ("M-Sj", AgeGroup.SUBJUNIOR),
            ("S-JR", AgeGroup.SUBJUNIOR), ("Sub-Junior", AgeGroup.SUBJUNIOR),
            ("JR", AgeGroup.JUNIOR), ("Junior", AgeGroup.JUNIOR),
            ("M-O", AgeGroup.OPEN), ("Open", AgeGroup.OPEN),
            ("M1", AgeGroup.M1), ("M-M2", AgeGroup.M2), ("M-M4", AgeGroup.M4),
        ):
            self.assertEqual(_division_label(code), expected, code)

    def test_a_missing_weight_class_comes_from_the_bodyweight(self):
        data = self._sheet(
            ["Място", "Име", "Пол", "Дивизия", "Лично тегло", "Тегл. кат.", "Най-доб.лег"],
            [[1, "Янек Кондев", "M", "Open", 81.2, "", 150]],
        )
        parsed = parse_upload("p.xlsx", data)
        ready, blocked = prepare_rows(
            parsed, default_sex="", default_equipment=Equipment.CLASSIC,
            default_event="auto", meet_date=date(2025, 5, 31),
        )
        self.assertEqual(blocked, [])
        self.assertEqual(ready[0].weight_class, "83")

    def test_a_missing_division_comes_from_the_age_in_years(self):
        data = self._sheet(
            ["Класиране", "Име", "Пол", "Възраст", "Дивизия", "Категория", "Най-доб.лег"],
            [[1, "Тест Тестов", "M", 45, "", "93", 150]],
        )
        parsed = parse_upload("p.xlsx", data)
        ready, _ = prepare_rows(
            parsed, default_sex="", default_equipment=Equipment.CLASSIC, default_event="auto"
        )
        self.assertEqual(ready[0].age_group, AgeGroup.M1)

    def test_a_club_points_line_is_not_a_lifter(self):
        data = self._sheet(
            ["Място", "Име", "Отбор", "Дивизия", "Категория", "Най-доб.лег"],
            [
                [1, "Истински Състезател", "НСА", "Open", "93", 150],
                [2, "Стренгт Скуад  12+7+12 =80", "", "", "", ""],
            ],
        )
        parsed = parse_upload("p.xlsx", data)
        self.assertEqual(len(parsed.rows), 1)
        self.assertEqual(parsed.skipped, 1)
        self.assertEqual(parsed.rows[0].raw_name, "Истински Състезател")

    def test_a_row_the_parser_could_not_read_is_reported_not_dropped(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["PL.", "Name", "1 Att.", "2 Att.", "3 Att.", "1 Att.", "2 Att.", "3 Att.", "RESULT"])
        sheet.append(["Open"])
        sheet.append(["- 93 kg"])
        sheet.append([1, "Ivanov Ivan", 100, 110, 120, 130, 140, 150, 270])
        parsed = parse_upload("odd.xlsx", _bytes(workbook))
        self.assertEqual(len(parsed.rows), 1)
        self.assertTrue(parsed.rows[0].errors)


class NameCaseTests(TestCase):
    def test_a_name_in_capitals_becomes_an_ordinary_name(self):
        self.assertEqual(normalize_name("ЕЛЕНА ЯНЕВА"), "Елена Янева")
        self.assertEqual(normalize_name("VALENTIN KOLEV 1"), "Valentin Kolev 1")
        self.assertEqual(normalize_name("ТАБАКОВА-ТРЕНДАФИЛОВА"), "Табакова-Трендафилова")

    def test_a_name_that_is_already_ordinary_is_left_alone(self):
        for name in ("Иван Петров", "Robert Mihaylov", "Ivan Petrov #2"):
            self.assertEqual(normalize_name(name), name)

    def test_extra_spacing_is_tidied(self):
        self.assertEqual(normalize_name("  Иван   Петров "), "Иван Петров")

    def test_a_digraph_inside_capitals_stays_capital(self):
        self.assertEqual(transliterate("ВЕНЦИСЛАВ"), "VENTSISLAV")
        self.assertEqual(transliterate("ЖИВКО ЩЕРЕВ"), "ZHIVKO SHTEREV")
        self.assertEqual(transliterate("Венцислав"), "Ventsislav")
        self.assertEqual(transliterate("Щерев"), "Shterev")

    def test_the_command_recases_and_rebuilds_the_latin_spelling(self):
        athlete = Athlete.objects.create(name_bg="ВЕНЦИСЛАВ КОСТАДИНОВ", sex=Sex.M)
        self.assertEqual(athlete.name_lat, "VENTSISLAV KOSTADINOV")
        call_command("fix_name_case", verbosity=0)
        athlete.refresh_from_db()
        self.assertEqual(athlete.name_bg, "Венцислав Костадинов")
        self.assertEqual(athlete.name_lat, "Ventsislav Kostadinov")

    def test_a_latin_name_from_its_own_source_is_only_recased(self):
        athlete = Athlete.objects.create(name_bg="ИВАН ПЕТРОВ", name_lat="PETROV IVAN", sex=Sex.M)
        call_command("fix_name_case", verbosity=0)
        athlete.refresh_from_db()
        self.assertEqual(athlete.name_bg, "Иван Петров")
        self.assertEqual(athlete.name_lat, "Petrov Ivan")

    def test_an_imported_name_in_capitals_is_stored_ordinary(self):
        from archive.services.commit import athlete_for

        athlete, _ = athlete_for("ГАЛИНА ИВАНОВА", Sex.F)
        self.assertEqual(athlete.name_bg, "Галина Иванова")


class BrokenDigraphRepairTests(TestCase):
    def test_a_latin_name_left_by_the_old_transliteration_is_rebuilt(self):
        athlete = Athlete.objects.create(name_bg="Тест", sex=Sex.M)
        Athlete.objects.filter(pk=athlete.pk).update(
            name_bg="ВЕНЦИСЛАВ КОСТАДИНОВ", name_lat="VENTsISLAV KOSTADINOV"
        )
        call_command("fix_name_case", verbosity=0)
        athlete.refresh_from_db()
        self.assertEqual(athlete.name_bg, "Венцислав Костадинов")
        self.assertEqual(athlete.name_lat, "Ventsislav Kostadinov")


class DuplicatesPageTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User

        self.client.force_login(User.objects.create_superuser("adm", "a@b.bg", "pw-for-tests-only"))
        self.one = Athlete.objects.create(name_bg="Александър Петров", name_lat="Aleksandar Petrov", sex=Sex.M)
        self.two = Athlete.objects.create(name_bg="Александър Петров", name_lat="Aleksander Petrov", sex=Sex.M)
        self.euro = Competition.objects.create(
            name="European Classic Powerlifting Championships", start_date=date(2021, 12, 3), slug="euro"
        )
        self.kard = Competition.objects.create(
            name="1 кръг, силов трибой с екип", start_date=date(2023, 4, 8), slug="kard"
        )
        self._start(self.one, self.kard, "66", Equipment.EQUIPPED, "1", 550)
        self._start(self.two, self.euro, "59", Equipment.CLASSIC, "2", Decimal("552.5"))

    def _start(self, athlete, competition, weight_class, equipment, place, total):
        return Result.objects.create(
            competition=competition, athlete=athlete, raw_name=athlete.name_bg, sex=Sex.M,
            age_group=AgeGroup.OPEN, equipment=equipment, event=Event.SBD,
            weight_class=weight_class, place=place, total=total, country="България",
        )

    def test_the_page_shows_each_start_with_its_meet_and_class(self):
        page = self.client.get("/admin/archive/athlete/duplicates/")
        self.assertContains(page, "European Classic Powerlifting Championships")
        self.assertContains(page, "1 кръг, силов трибой с екип")
        self.assertContains(page, ">66<")
        self.assertContains(page, ">59<")
        self.assertContains(page, "Слей тези двама")

    def test_two_lifters_at_one_meet_are_flagged_and_cannot_be_merged_from_here(self):
        self._start(self.two, self.kard, "93", Equipment.EQUIPPED, "4", 600)
        page = self.client.get("/admin/archive/athlete/duplicates/")
        self.assertContains(page, "значи са различни хора")
        self.assertNotContains(page, "Слей тези двама")

    def test_the_athlete_page_lists_the_starts(self):
        page = self.client.get(f"/admin/archive/athlete/{self.two.pk}/change/")
        self.assertContains(page, "Стартове")
        self.assertContains(page, "European Classic Powerlifting Championships")


class DivisionCodeShapeTests(TestCase):
    def test_the_shapes_the_sources_use(self):
        from archive.services.importers import _division_label

        for code, expected in (
            ("F-C-Open", AgeGroup.OPEN), ("M-C-Open", AgeGroup.OPEN),
            ("MR-O", AgeGroup.OPEN), ("FR-Jr", AgeGroup.JUNIOR),
            ("MR-Sj", AgeGroup.SUBJUNIOR), ("MR-M1", AgeGroup.M1),
            ("Sub-Junior", AgeGroup.SUBJUNIOR), ("Masters 1", AgeGroup.M1),
        ):
            self.assertEqual(_division_label(code), expected, code)

    def test_a_heading_is_not_mistaken_for_a_code(self):
        from archive.services.importers import _division_label

        for heading in ("Best Lifters of Subjuniors", "Nation (points)", "Best Lifters of Seniors"):
            self.assertEqual(_division_label(heading), "", heading)


class SplitNameAndAgeClassTests(TestCase):
    def _sheet(self, header, row):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(header)
        sheet.append(row)
        return parse_upload("p.xlsx", _bytes(workbook))

    def test_a_name_split_over_two_columns_is_joined(self):
        parsed = self._sheet(
            ["№", "Име", "Фамилия", "Отбор", "Дивизия", "кат.", "лег"],
            [1, "Станимир", "Данчев", "Марек", "M-CL-BP", "59", 75],
        )
        self.assertEqual(parsed.rows[0].raw_name, "Станимир Данчев")
        self.assertEqual(parsed.rows[0].club, "Марек")

    def test_the_age_class_column_gives_the_division(self):
        parsed = self._sheet(
            ["№", "Име", "Фамилия", "ГР", "Дивизия", "кат.", "лег"],
            [1, "Станимир", "Данчев", 18, "M-CL-BP", "59", 75],
        )
        ready, blocked = prepare_rows(
            parsed, default_sex="", default_equipment="auto", default_event="auto"
        )
        self.assertEqual(blocked, [])
        self.assertEqual(ready[0].age_group, AgeGroup.SUBJUNIOR)
        self.assertEqual(ready[0].event, Event.B)
        self.assertEqual(ready[0].equipment, Equipment.CLASSIC)

    def test_the_age_class_numbers_the_protocols_use(self):
        for value, expected in ((18, AgeGroup.SUBJUNIOR), (23, AgeGroup.JUNIOR),
                                (24, AgeGroup.OPEN), (40, AgeGroup.M1), (50, AgeGroup.M2)):
            parsed = self._sheet(
                ["№", "Име", "Фамилия", "ГР", "Дивизия", "кат.", "лег"],
                [1, "Тест", "Тестов", value, "M-CL-BP", "93", 100],
            )
            ready, _ = prepare_rows(
                parsed, default_sex="", default_equipment="auto", default_event="auto"
            )
            self.assertEqual(ready[0].age_group, expected, value)

    def test_a_birth_year_in_that_column_is_ignored(self):
        parsed = self._sheet(
            ["№", "Име", "Фамилия", "ГР", "Дивизия", "кат.", "лег"],
            [1, "Тест", "Тестов", 1995, "M-CL-BP", "93", 100],
        )
        _, blocked = prepare_rows(
            parsed, default_sex="", default_equipment="auto", default_event="auto"
        )
        self.assertIn("Няма възрастова група", blocked[0]["problems"][0])


class TickMarkAttemptTests(TestCase):
    def _row(self, cells):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["№", "Име", "Фамилия", "ГР", "Дивизия", "тегло", "кат.",
                      "лег-1", "", "лег-2", "", "лег-3", "", "лег"])
        sheet.append(cells)
        return parse_upload("p.xlsx", _bytes(workbook)).rows[0]

    def test_the_weight_sits_beside_the_tick(self):
        row = self._row([1, "Станимир", "Данчев", 18, "M-CL-BP", 56.7, "59",
                         "√", 60, "√", 70, "√", 75, 75])
        self.assertEqual(row.attempts["bench1"], 60)
        self.assertEqual(row.attempts["bench2"], 70)
        self.assertEqual(row.attempts["bench3"], 75)
        self.assertEqual(row.best_bench, 75)

    def test_a_cross_marks_the_attempt_failed(self):
        row = self._row([1, "Тест", "Тестов", 24, "M-CL-BP", 82.0, "83",
                         "√", 100, "×", 110, "×", 110, 100])
        self.assertEqual(row.attempts["bench1"], 100)
        self.assertEqual(row.attempts["bench2"], -110)
        self.assertEqual(row.attempts["bench3"], -110)
        self.assertEqual(row.best_bench, 100)

    def test_a_plain_x_still_means_a_failed_attempt_with_no_weight(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["PL.", "Name", "Nation", "Weight", "1 Att.", "2 Att.", "3 Att.", "RESULT"])
        sheet.append(["Open"])
        sheet.append(["- 93 kg"])
        sheet.append([1, "Ivanov Ivan", "NSA", 92.5, 100, 110, "X", 110])
        row = parse_upload("g.xlsx", _bytes(workbook)).rows[0]
        self.assertIsNone(row.attempts.get("bench3"))
        self.assertEqual(row.best_bench, 110)


class ProtocolReplacesOplSourceTests(TestCase):
    def test_loading_a_protocol_drops_the_openpowerlifting_note(self):
        competition = Competition.objects.create(
            name="Вдигане от лег", start_date=date(2024, 9, 21), slug="dupn-2024"
        )
        CompetitionFile.objects.create(
            competition=competition, kind=FileKind.OPL,
            url="https://www.openpowerlifting.org/m/bulgarianpf/2402",
        )
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["№", "Име", "Фамилия", "ГР", "Дивизия", "кат.", "лег"])
        sheet.append([1, "Станимир", "Данчев", 18, "M-CL-BP", "59", 75])
        path = Path(self._tmp()) / "p.xlsx"
        workbook.save(path)
        call_command("import_protocol", str(path), competition="dupn-2024", verbosity=0)
        self.assertEqual(competition.results.count(), 1)
        self.assertFalse(competition.files.filter(kind=FileKind.OPL).exists())

    def test_adding_a_second_file_keeps_it(self):
        competition = Competition.objects.create(
            name="Лег", start_date=date(2026, 5, 30), slug="sofia-leg"
        )
        CompetitionFile.objects.create(
            competition=competition, kind=FileKind.OPL, url="https://example.org/m/1"
        )
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["№", "Име", "Фамилия", "ГР", "Дивизия", "кат.", "лег"])
        sheet.append([1, "Тест", "Тестов", 24, "M-CL-BP", "93", 100])
        path = Path(self._tmp()) / "q.xlsx"
        workbook.save(path)
        call_command("import_protocol", str(path), competition="sofia-leg", keep=True, verbosity=0)
        self.assertTrue(competition.files.filter(kind=FileKind.OPL).exists())

    def _tmp(self):
        import tempfile

        directory = tempfile.mkdtemp()
        self.addCleanup(__import__("shutil").rmtree, directory, True)
        return directory


class PlaceColumnTests(TestCase):
    def _sheet(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["№", "№", "Име", "Фамилия", "Отбор", "ГР", "Дивизия", "кат.", "лег"])
        for row in rows:
            sheet.append(row)
        return parse_upload("p.xlsx", _bytes(workbook))

    def test_the_placing_is_taken_over_the_running_number(self):
        parsed = self._sheet([
            [1, 1, "Първи", "Един", "НСА", 24, "M-CL-BP", "93", 150],
            [2, 2, "Втори", "Два", "НСА", 24, "M-CL-BP", "93", 140],
            [3, 1, "Трети", "Три", "НСА", 24, "M-CL-BP", "105", 200],
            [4, 2, "Четвърти", "Четири", "НСА", 24, "M-CL-BP", "105", 190],
        ])
        self.assertEqual([row.place for row in parsed.rows], ["1", "2", "1", "2"])

    def test_a_single_number_column_is_still_the_place(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Място", "Име", "Дивизия", "кат.", "Най-доб.лег"])
        sheet.append([3, "Един Единов", "Open", "93", 150])
        parsed = parse_upload("p.xlsx", _bytes(workbook))
        self.assertEqual(parsed.rows[0].place, "3")

    def test_the_running_number_is_not_used_when_it_never_repeats(self):
        parsed = self._sheet([[i, 1, f"Име{i}", f"Фам{i}", "НСА", 24, "M-CL-BP", str(50 + i), 100]
                              for i in range(1, 7)])
        self.assertEqual({row.place for row in parsed.rows}, {"1"})


class MeetPageWidthTests(TestCase):
    def test_the_meet_page_is_marked_for_the_full_width_layout(self):
        Competition.objects.create(name="Турнир", start_date=date(2025, 5, 1), slug="w")
        page = self.client.get("/competitions/w/")
        self.assertContains(page, '<body class="meet">')

    def test_the_reading_pages_keep_the_centred_column(self):
        for path in ("/", "/competitions/", "/athletes/", "/records/"):
            self.assertContains(self.client.get(path), '<body class="">', msg_prefix=path)


class ColumnBandingTests(TestCase):
    def setUp(self):
        self.competition = Competition.objects.create(
            name="Турнир", start_date=date(2025, 5, 1), slug="band"
        )
        athlete = Athlete.objects.create(name_bg="Тест Тестов", sex=Sex.M)
        Result.objects.create(
            competition=self.competition, athlete=athlete, raw_name="Тест Тестов", sex=Sex.M,
            age_group=AgeGroup.OPEN, equipment=Equipment.CLASSIC, event=Event.SBD,
            weight_class="93", country="България", squat1=200, bench1=150, deadlift1=250,
            best_squat=200, best_bench=150, best_deadlift=250, total=600,
        )

    def test_each_lift_and_the_total_carry_their_own_class(self):
        page = self.client.get("/competitions/band/").content.decode()
        for name in ("lift-squat", "lift-bench", "lift-deadlift", "col-total"):
            self.assertIn(name, page, name)
        # three attempt cells plus the heading that spans them
        self.assertEqual(page.count('class="lift-bench"'), 4)

    def test_a_bench_only_meet_bands_nothing_it_did_not_contest(self):
        self.competition.results.update(
            event=Event.B, best_squat=None, best_deadlift=None, squat1=None,
            deadlift1=None, total=None,
        )
        page = self.client.get("/competitions/band/").content.decode()
        self.assertIn("lift-bench", page)
        self.assertNotIn("lift-squat", page)
        self.assertNotIn("col-total", page)

    def test_the_athlete_page_bands_the_same_way(self):
        athlete = Athlete.objects.get(name_bg="Тест Тестов")
        page = self.client.get(f"/athletes/{athlete.slug}/").content.decode()
        for name in ("lift-squat", "lift-bench", "lift-deadlift", "col-total"):
            self.assertIn(name, page, name)


class FilterFormLayoutTests(TestCase):
    def test_the_records_filter_and_its_button_share_one_row(self):
        page = self.client.get("/records/").content.decode()
        self.assertIn('<form class="filters" method="get">', page)
        self.assertIn("<p>", page)  # Django's as_p wraps each field

    def test_the_stylesheet_flattens_those_paragraphs(self):
        css = (Path("archive/static/archive/site.css")).read_text(encoding="utf-8")
        self.assertIn(".filters p { margin: 0; }", css)
